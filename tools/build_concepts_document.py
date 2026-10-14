from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUTPUT_PATH = r"C:\Training2026\OrderFlow\OrderFlow_Implemented_Concepts_Guide.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text_color(cell, color):
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor.from_string(color)


def set_cell_borders(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def style_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(4)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
            if row_index == 0:
                set_cell_shading(cell, "1F4E79")
                set_cell_text_color(cell, "FFFFFF")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F4F8FB")


def add_intro(doc):
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("OrderFlow Implemented Concepts Guide")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Definitions and exact code locations for Phases 1 through 17")
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(80, 80, 80)

    doc.add_paragraph(
        "This document explains the important concepts implemented in the OrderFlow learning project so far. "
        "Each section defines the concept in beginner-friendly language, explains why it exists in an enterprise application, "
        "and identifies where it appears in this codebase."
    )

    doc.add_paragraph(
        "The current solution contains the original ASP.NET Core API for early-phase identity and cart learning, "
        "six extracted microservices, a YARP gateway, React and Angular portals, RabbitMQ messaging, Redis caching, "
        "OpenTelemetry tracing, and a .NET Aspire AppHost for local orchestration."
    )


def add_architecture_snapshot(doc):
    doc.add_heading("Architecture Snapshot", level=1)
    table = doc.add_table(rows=1, cols=3)
    hdr = table.rows[0].cells
    hdr[0].text = "Area"
    hdr[1].text = "Purpose"
    hdr[2].text = "Main location"
    rows = [
        ("Original API", "Identity, JWT login, earlier monolith endpoints, cart and compatibility learning.", r"C:\Training2026\OrderFlow\OrderFlow.Api"),
        ("Microservices", "Separate customer, product, inventory, order, payment, and notification services.", r"C:\Training2026\OrderFlow\Services"),
        ("Gateway", "Single entry point for portal-to-service traffic using YARP.", r"C:\Training2026\OrderFlow\Gateway"),
        ("Customer Portal", "React app for customer shopping and checkout flow.", r"C:\Training2026\OrderFlow\CustomerPortal"),
        ("Admin Portal", "Angular app for admin product, customer, inventory, and order management.", r"C:\Training2026\OrderFlow\AdminPortal"),
        ("Aspire AppHost", "Local orchestration of services, gateway, Redis, RabbitMQ, and original API.", r"C:\Training2026\OrderFlow\OrderFlow.AppHost"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
    style_table(table)


def add_concept(doc, title, definition, why, locations):
    doc.add_heading(title, level=2)
    p = doc.add_paragraph()
    p.add_run("Definition: ").bold = True
    p.add_run(definition)
    p = doc.add_paragraph()
    p.add_run("Why it matters: ").bold = True
    p.add_run(why)

    doc.add_paragraph("Implemented in:", style=None).runs[0].bold = True
    for path, look_for, note in locations:
        item = doc.add_paragraph(style="List Bullet")
        item.add_run(path).bold = True
        item.add_run(f" - {look_for} {note}")


def add_concepts(doc):
    doc.add_heading("Backend Concepts", level=1)
    concepts = [
        (
            "ASP.NET Core Web API",
            "A backend HTTP application that exposes endpoints for browser clients or other services.",
            "It gives the portals and services a standard way to request data and trigger business actions.",
            [
                (r"C:\Training2026\OrderFlow\OrderFlow.Api\Program.cs", "App startup, dependency injection, middleware, Swagger, authentication, and database setup.", "This is the original API used for auth and earlier-phase modules."),
                (r"C:\Training2026\OrderFlow\Services\*\Program.cs", "Each microservice has its own startup pipeline.", "Enterprise systems usually repeat this startup pattern per service."),
            ],
        ),
        (
            "Controller Based API Structure",
            "A controller groups related HTTP endpoints into a class such as ProductsController or OrdersController.",
            "It keeps endpoint code discoverable and maps naturally to enterprise feature ownership.",
            [
                (r"C:\Training2026\OrderFlow\Services\ProductService\Controllers\ProductsController.cs", "Product CRUD endpoints.", "Product behavior lives in the product service, not in the portal."),
                (r"C:\Training2026\OrderFlow\Services\OrderService\Controllers\CheckoutController.cs", "Checkout workflow endpoint and saga coordination.", "This controller coordinates several services but still keeps service boundaries visible."),
                (r"C:\Training2026\OrderFlow\Services\InventoryService\Controllers\InventoryController.cs", "Inventory endpoints for stock management and reservation.", "Inventory is isolated from products so stock can evolve independently."),
            ],
        ),
        (
            "Entity Model",
            "An entity is a C# class that represents business data stored in the database.",
            "Entities are the durable shape of the business state, such as Product, Customer, Order, and InventoryItem.",
            [
                (r"C:\Training2026\OrderFlow\Services\ProductService\Models\Product.cs", "Product database model.", "Represents catalog data."),
                (r"C:\Training2026\OrderFlow\Services\OrderService\Models\Order.cs", "Order and CheckoutSagaState models.", "Represents order state and the saga state machine."),
                (r"C:\Training2026\OrderFlow\Services\NotificationService\Models\Notification.cs", "Notification model with SourceEventId.", "SourceEventId supports idempotent event handling."),
            ],
        ),
        (
            "DTO",
            "A Data Transfer Object is a class or record used to receive or return API data without exposing the full database entity.",
            "DTOs protect the internal model and make API contracts clear.",
            [
                (r"C:\Training2026\OrderFlow\Services\OrderService\Dtos\OrderDtos.cs", "CreateOrderRequest, CheckoutRequest, CheckoutResponse, and related records.", "The checkout contract is separate from the database model."),
                (r"C:\Training2026\OrderFlow\Services\ProductService\Dtos\ProductDtos.cs", "Product create/update request contracts.", "The browser sends DTOs, not EF entities."),
            ],
        ),
        (
            "Entity Framework Core DbContext",
            "DbContext is the EF Core unit that maps C# entities to SQL Server tables and performs database operations.",
            "It lets services work with typed C# objects while EF handles SQL generation.",
            [
                (r"C:\Training2026\OrderFlow\Services\ProductService\Data\ProductDbContext.cs", "Products DbSet and database mapping.", "Each service has its own database context."),
                (r"C:\Training2026\OrderFlow\Services\OrderService\Data\OrderDbContext.cs", "Orders, OrderItems, and CheckoutSagas mappings.", "Order persistence is separate from Product persistence."),
                (r"C:\Training2026\OrderFlow\OrderFlow.Api\Data\OrderFlowDbContext.cs", "Original monolith database context.", "Still used by the original API for earlier-phase auth/cart learning."),
            ],
        ),
        (
            "ASP.NET Core Identity and JWT",
            "Identity stores users and roles; JWT is a signed token sent by the client after login.",
            "This allows the admin portal to log in once and call protected APIs without sending the password on every request.",
            [
                (r"C:\Training2026\OrderFlow\OrderFlow.Api\Program.cs", "Identity, authentication, authorization, JWT bearer configuration, and seed users.", "The admin login depends on this API."),
                (r"C:\Training2026\OrderFlow\AdminPortal\src\app\features\auth\auth-api.service.ts", "Login request and token storage.", "The Angular app stores the JWT after login."),
                (r"C:\Training2026\OrderFlow\AdminPortal\src\app\features\auth\auth.interceptor.ts", "Bearer token for original API requests.", "The interceptor attaches auth headers automatically."),
            ],
        ),
        (
            "Microservices",
            "A microservice is a separately running application responsible for one business capability.",
            "It helps you learn service boundaries, separate databases, independent deployments, and inter-service communication.",
            [
                (r"C:\Training2026\OrderFlow\Services\CustomerService", "Customer management service.", "Owns customer data."),
                (r"C:\Training2026\OrderFlow\Services\ProductService", "Product catalog service.", "Owns product data and product caching."),
                (r"C:\Training2026\OrderFlow\Services\OrderService", "Order and checkout service.", "Coordinates checkout and stores order/saga state."),
                (r"C:\Training2026\OrderFlow\Services\InventoryService", "Inventory service.", "Owns stock and reservation behavior."),
                (r"C:\Training2026\OrderFlow\Services\PaymentService", "Mock payment service.", "Simulates payment success or failure."),
                (r"C:\Training2026\OrderFlow\Services\NotificationService", "Notification service.", "Consumes events and stores notifications."),
            ],
        ),
        (
            "Synchronous Service Communication",
            "One service calls another service directly over HTTP and waits for the response.",
            "Checkout needs current product, inventory, and payment responses before it can decide the final order status.",
            [
                (r"C:\Training2026\OrderFlow\Services\OrderService\Program.cs", "Named HttpClient registrations and resilience handlers.", "OrderService knows the downstream service base URLs."),
                (r"C:\Training2026\OrderFlow\Services\OrderService\Controllers\CheckoutController.cs", "HTTP calls to product, inventory, payment, and notification services.", "This is the main synchronous workflow."),
            ],
        ),
        (
            "RabbitMQ Messaging",
            "RabbitMQ is a message broker. Services publish events to it, and other services consume those events asynchronously.",
            "Messaging reduces direct coupling and lets background processes react to business events.",
            [
                (r"C:\Training2026\OrderFlow\Services\OrderService\Messaging\RabbitMqEventPublisher.cs", "Publishes integration events to the RabbitMQ exchange.", "OrderService does not directly know every future event consumer."),
                (r"C:\Training2026\OrderFlow\Services\OrderService\Messaging\IntegrationEvent.cs", "Shared event envelope shape.", "EventId supports tracking and idempotency."),
                (r"C:\Training2026\OrderFlow\Services\NotificationService\Messaging\RabbitMqNotificationConsumer.cs", "Consumes order events from RabbitMQ.", "NotificationService reacts in the background."),
            ],
        ),
        (
            "Integration Event",
            "An integration event is a message that says something important happened in one service.",
            "Events let other services react without turning the original action into one large direct call chain.",
            [
                (r"C:\Training2026\OrderFlow\Services\OrderService\Messaging\IntegrationEvent.cs", "EventId, EventType, OrderId, OccurredAtUtc, and Payload.", "This is the envelope for events like order.confirmed and payment.failed."),
                (r"C:\Training2026\OrderFlow\Services\OrderService\Controllers\CheckoutController.cs", "Publishes inventory.reserved, order.confirmed, and payment.failed events.", "Checkout emits business facts as it progresses."),
            ],
        ),
        (
            "Dead Letter Queue",
            "A dead letter queue stores messages that could not be processed successfully after retries.",
            "It prevents one bad message from blocking the queue forever while preserving it for later diagnosis.",
            [
                (r"C:\Training2026\OrderFlow\Services\NotificationService\Messaging\RabbitMqNotificationConsumer.cs", "Queue declaration, retry count, and dead-letter behavior.", "Failed messages move to the configured DLQ after retry attempts."),
                (r"C:\Training2026\OrderFlow\Services\NotificationService\appsettings.json", "RabbitMQ queue and DLQ names.", "Queue names are configuration, not code constants."),
            ],
        ),
        (
            "Idempotency",
            "Idempotency means repeating the same request or message does not create duplicate business effects.",
            "Distributed systems retry work, so code must safely handle duplicates.",
            [
                (r"C:\Training2026\OrderFlow\Services\OrderService\Models\Order.cs", "CheckoutSagaState contains IdempotencyKey.", "Duplicate checkout requests can be detected."),
                (r"C:\Training2026\OrderFlow\Services\OrderService\Data\OrderDbContext.cs", "Unique index on saga IdempotencyKey.", "The database enforces uniqueness."),
                (r"C:\Training2026\OrderFlow\Services\NotificationService\Models\Notification.cs", "SourceEventId property.", "Duplicate events do not create duplicate notifications."),
                (r"C:\Training2026\OrderFlow\Services\NotificationService\Data\NotificationDbContext.cs", "Unique filtered index on SourceEventId.", "Idempotency is protected at the database level."),
            ],
        ),
        (
            "Saga Pattern",
            "A saga is a workflow that coordinates multiple services using steps and compensating actions instead of one database transaction.",
            "Checkout touches product, inventory, payment, order, and notification boundaries; a normal SQL transaction cannot cover all of them.",
            [
                (r"C:\Training2026\OrderFlow\Services\OrderService\Controllers\CheckoutController.cs", "Saga steps, failure handling, and compensation.", "This is the heart of distributed checkout."),
                (r"C:\Training2026\OrderFlow\Services\OrderService\Models\Order.cs", "CheckoutSagaState model.", "The saga has durable state."),
                (r"C:\Training2026\OrderFlow\Services\OrderService\Data\OrderDbContext.cs", "CheckoutSagas table mapping.", "Saga progress is saved in SQL Server."),
            ],
        ),
        (
            "Compensating Action",
            "A compensating action reverses or offsets a previous successful step when a later step fails.",
            "If payment fails after inventory is reserved, the inventory must be released.",
            [
                (r"C:\Training2026\OrderFlow\Services\OrderService\Controllers\CheckoutController.cs", "Failure path calls inventory release and cancels the order.", "This replaces a cross-service transaction."),
                (r"C:\Training2026\OrderFlow\Services\InventoryService\Controllers\InventoryController.cs", "Inventory reserve and release endpoints.", "Inventory provides the operations needed by the saga."),
            ],
        ),
        (
            "Resilience",
            "Resilience is the ability to survive temporary failures such as slow services, network errors, or overloaded endpoints.",
            "Microservices fail in partial ways; resilience policies make those failures controlled and visible.",
            [
                (r"C:\Training2026\OrderFlow\Services\OrderService\Program.cs", "Standard resilience handler on HttpClient registrations.", "Adds retry, timeout, and circuit-breaker style behavior."),
                (r"C:\Training2026\OrderFlow\Gateway\Program.cs", "Rate limiter and health checks.", "The gateway protects downstream services from too much traffic."),
            ],
        ),
        (
            "Health Check",
            "A health check is a lightweight endpoint that reports whether an app is alive and ready enough to be monitored.",
            "Gateways, orchestrators, and developers need a quick signal that a service is running.",
            [
                (r"C:\Training2026\OrderFlow\Services\*\Program.cs", "AddHealthChecks and MapHealthChecks('/health').", "Each service exposes a health endpoint."),
                (r"C:\Training2026\OrderFlow\Gateway\Program.cs", "Gateway health endpoint.", "You can test the gateway at http://localhost:6200/health."),
            ],
        ),
        (
            "Redis Cache Aside",
            "Cache-aside means the application checks the cache first; on a miss, it loads from the database and stores the result in cache.",
            "Product reads are common, so caching reduces database load and improves response time.",
            [
                (r"C:\Training2026\OrderFlow\Services\ProductService\Controllers\ProductsController.cs", "Product list cache lookup, database fallback, expiration, and invalidation.", "This is where GET /api/products becomes cache-aware."),
                (r"C:\Training2026\OrderFlow\Services\ProductService\Program.cs", "Redis connection registration.", "The controller receives Redis through dependency injection."),
                (r"C:\Training2026\OrderFlow\Services\ProductService\appsettings.json", "Redis connection string.", "Local default is localhost:6379 unless Aspire supplies it."),
            ],
        ),
        (
            "YARP API Gateway",
            "YARP is Microsoft's reverse proxy library for routing incoming requests to backend services.",
            "The portals get one stable API entry point while services stay independently hosted behind it.",
            [
                (r"C:\Training2026\OrderFlow\Gateway\Program.cs", "Reverse proxy, CORS, rate limiting, health checks, authentication, and tracing.", "The gateway contains infrastructure logic, not business logic."),
                (r"C:\Training2026\OrderFlow\Gateway\appsettings.json", "Routes and clusters mapping URLs to service ports.", "This is where /api/products maps to ProductService."),
                (r"C:\Training2026\OrderFlow\Gateway\Auth\BasicAuthenticationHandler.cs", "Development Basic authentication for proxied routes.", "Covers the gateway auth integration requirement."),
            ],
        ),
        (
            "CORS",
            "Cross Origin Resource Sharing allows browser apps on one port to call an API on another port.",
            "React, Angular, and the gateway all run on different localhost ports during development.",
            [
                (r"C:\Training2026\OrderFlow\Gateway\Program.cs", "CORS policy for localhost:5173 and localhost:4200.", "Allows both portals to call the gateway."),
                (r"C:\Training2026\OrderFlow\OrderFlow.Api\Program.cs", "CORS for earlier-phase API usage.", "Supports auth and original API calls."),
            ],
        ),
        (
            "OpenTelemetry",
            "OpenTelemetry captures traces and metrics so you can understand how a request moves through distributed components.",
            "In microservices, one browser action can involve several services; tracing connects those hops.",
            [
                (r"C:\Training2026\OrderFlow\Gateway\Program.cs", "Gateway tracing registration and console exporter.", "Shows requests entering the system."),
                (r"C:\Training2026\OrderFlow\Services\OrderService\Program.cs", "OrderService tracing and HTTP client instrumentation.", "Shows checkout calls to downstream services."),
            ],
        ),
        (
            ".NET Aspire AppHost",
            "Aspire AppHost is a local orchestration project that starts and connects multiple app resources for development.",
            "It makes the distributed project easier to run without manually opening many terminals.",
            [
                (r"C:\Training2026\OrderFlow\OrderFlow.AppHost\Program.cs", "Registers Redis, RabbitMQ, gateway, services, and original API.", "This is the local distributed application map."),
                (r"C:\Training2026\OrderFlow\OrderFlow.AppHost\OrderFlow.AppHost.csproj", "Aspire hosting package references and project references.", "AppHost knows which projects belong to the local system."),
            ],
        ),
        (
            "React Customer Portal",
            "React is the frontend library used for the customer shopping experience.",
            "It lets customers browse products, manage cart state, and perform checkout.",
            [
                (r"C:\Training2026\OrderFlow\CustomerPortal\src\shared\api.ts", "Gateway base URLs and Basic auth header for gateway requests.", "The portal talks to microservices through the gateway."),
                (r"C:\Training2026\OrderFlow\CustomerPortal\src\features", "Feature folders for customer, product, order, and cart concepts.", "This mirrors enterprise frontend organization."),
            ],
        ),
        (
            "Angular Admin Portal",
            "Angular is the frontend framework used for the admin management experience.",
            "It provides structured services, components, routing, and interceptors for admin screens.",
            [
                (r"C:\Training2026\OrderFlow\AdminPortal\src\app\features\products\product-api.service.ts", "Product API service calling the gateway.", "Feature-specific API service."),
                (r"C:\Training2026\OrderFlow\AdminPortal\src\app\features\inventory\inventory-api.service.ts", "Inventory API service calling the gateway.", "Inventory UI talks to InventoryService through YARP."),
                (r"C:\Training2026\OrderFlow\AdminPortal\src\app\features\auth\auth.interceptor.ts", "Bearer token for original API, Basic header for gateway.", "Shows how cross-cutting frontend auth is centralized."),
            ],
        ),
        (
            "Configuration",
            "Configuration stores values that change between environments, such as ports, connection strings, queues, and service URLs.",
            "Enterprise apps avoid hard-coding infrastructure details inside business logic.",
            [
                (r"C:\Training2026\OrderFlow\Services\OrderService\appsettings.json", "Service URLs and RabbitMQ settings.", "OrderService can be pointed at different downstream endpoints."),
                (r"C:\Training2026\OrderFlow\Gateway\appsettings.json", "Gateway routes, clusters, and Basic auth defaults.", "Routing can change without rewriting controllers."),
                (r"C:\Training2026\OrderFlow\Services\ProductService\appsettings.json", "SQL Server and Redis connection strings.", "ProductService owns its data and cache configuration."),
            ],
        ),
    ]

    for concept in concepts:
        add_concept(doc, *concept)


def add_testing_section(doc):
    doc.add_heading("How to Test the Implemented Concepts Locally", level=1)
    doc.add_paragraph(
        "For Phase 17, the simplest test path is to run the Aspire AppHost, then run both portals. "
        "Aspire starts the backend pieces, while Vite and Angular serve the two browser applications."
    )
    table = doc.add_table(rows=1, cols=3)
    hdr = table.rows[0].cells
    hdr[0].text = "Step"
    hdr[1].text = "Command or URL"
    hdr[2].text = "What it verifies"
    rows = [
        ("Start backend", r"dotnet run --project C:\Training2026\OrderFlow\OrderFlow.AppHost\OrderFlow.AppHost.csproj", "Aspire, services, gateway, Redis, RabbitMQ, and original API startup."),
        ("Start customer portal", r"cd C:\Training2026\OrderFlow\CustomerPortal && npm run dev", "React portal loads and calls gateway routes."),
        ("Start admin portal", r"cd C:\Training2026\OrderFlow\AdminPortal && npm start", "Angular portal loads, logs in through original API, and manages service data through gateway."),
        ("Gateway health", "http://localhost:6200/health", "YARP gateway process is reachable."),
        ("Admin login", "admin@oms.local / ChangeMe123!", "Identity and JWT flow still work through OrderFlow.Api."),
        ("Checkout flow", "Create products and inventory, add to customer cart, then checkout.", "Synchronous service calls, saga, inventory compensation, payment, and RabbitMQ events."),
    ]
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
    style_table(table)


def configure_styles(doc):
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    for style_name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.color.rgb = RGBColor(0, 0, 0)
    styles["Heading 1"].font.size = Pt(18)
    styles["Heading 2"].font.size = Pt(14)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    configure_styles(doc)
    add_intro(doc)
    add_architecture_snapshot(doc)
    add_concepts(doc)
    add_testing_section(doc)
    doc.save(OUTPUT_PATH)


if __name__ == "__main__":
    build()
